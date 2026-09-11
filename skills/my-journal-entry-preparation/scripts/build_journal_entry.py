#!/usr/bin/env python3
"""
Build a SCALABLE journal-entry workbook (.xlsx) from a JSON description.

Designed for tens to thousands of entries: the master register is one compact row
per entry (not a block per entry), with a normalized Journal Lines sheet for the
debit/credit breakdown, and a wide Entry Details sheet (still one row per entry)
for the full audit trail. Every sheet is linked by Entry ID, with hyperlinks from
the register to the matching Entry Details row and back.

Usage:
    python build_journal_entry.py --input entries.json --output journal_entries.xlsx

See EXAMPLE_JSON at the bottom of this file for a fully worked example of the
input shape. Key top-level keys in the input JSON:

  "entries"           - list of entry dicts (see below)
  "exceptions"         - list of exception dicts (optional)
  "source_documents"   - list of source-document dicts (optional)

Each entry dict:
  entry_id, status, entry_type, transaction_date, posting_date, accounting_period,
  description, source_reference, currency, narration, calculation_method,
  assumptions (list[str]), confidence ("High"/"Medium"/"Low"),
  confidence_details (list of {field, level, reason}), tax_treatment,
  reversal_required (bool), reversal_date, preparer, reviewer, approver,
  created_at, reviewed_at, approved_at, posted_at, change_history (list[str]),
  rejection_reason, correction_reference,
  lines: list of {line_no, account_code, account_name, debit, credit, narration,
                   vendor_customer, department, cost_centre, project, tax_code,
                   source_doc_ref}

Each exception dict:
  entry_id, exception_type, risk_level ("High"/"Medium"/"Low"), description,
  required_action, assigned_reviewer, resolution_status, resolution_notes

Each source_document dict:
  entry_id, document_id, document_type, doc_number, vendor_customer,
  document_date, document_amount, original_filename, file_link,
  ocr_confidence, verification_status
"""

import argparse
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink

FONT_NAME = "Arial"
MAX_ROW = 20000  # formula range ceiling -- generous for "hundreds, or thousands" of entries

HEADER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color="1F4E5F")
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="777777")
KPI_LABEL_FONT = Font(name=FONT_NAME, size=10, color="555555")
KPI_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=18, color="1F4E5F")
BOLD_FONT = Font(name=FONT_NAME, bold=True)
WARN_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
BAD_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
GOOD_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
NEUTRAL_FILL = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
KPI_FILL = PatternFill(start_color="EAF1F4", end_color="EAF1F4", fill_type="solid")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY_FMT = '#,##0.00;(#,##0.00);"-"'
HYPERLINK_FONT = Font(name=FONT_NAME, color="1155CC", underline="single")

STATUS_FILL = {
    "Draft": NEUTRAL_FILL,
    "Pending Review": WARN_FILL,
    "Under Review": WARN_FILL,
    "Approved": GOOD_FILL,
    "Posted": GOOD_FILL,
    "Rejected": BAD_FILL,
    "Reversed": BAD_FILL,
    "Corrected": WARN_FILL,
}
LEVEL_FILL = {"High": GOOD_FILL, "Medium": WARN_FILL, "Low": BAD_FILL}


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_table(ws, name, ncols, nrows, header_row=1):
    """Add an Excel Table (gives filter dropdowns + sortable columns + banded rows)."""
    if nrows < 1:
        return
    last_col = get_column_letter(ncols)
    last_row = header_row + nrows
    ref = f"A{header_row}:{last_col}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def link_to(cell, sheet_name, cell_ref):
    """Set a proper IN-WORKBOOK hyperlink (uses the 'location' field, not 'target' --
    assigning a '#...' string to cell.hyperlink writes it as an external target and gets
    percent-encoded, which breaks the link)."""
    cell.hyperlink = Hyperlink(ref="", location=f"'{sheet_name}'!{cell_ref}")
    cell.font = HYPERLINK_FONT


CLIENT_BANNER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
CLIENT_BANNER_FONT = Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")


def add_client_banner(ws, row, client, span_cols=10):
    """Prominent 'Active Client: CLI-0001 -- ABC Pvt Ltd' banner, placed on every sheet so the
    active client is never ambiguous. Uses a row that's otherwise blank in that sheet's layout
    (chosen per-sheet by the caller) so it never collides with existing headers/formulas."""
    label = f"Active Client: {client.get('client_id', '')} -- {client.get('legal_name', '')}"
    sub = " | ".join(
        v for v in [
            client.get("financial_year", "") and f"FY {client['financial_year']}",
            client.get("base_currency", ""),
            client.get("jurisdiction", "") or client.get("country", ""),
        ] if v
    )
    if sub:
        label = f"{label}   ({sub})"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span_cols)
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = CLIENT_BANNER_FONT
    cell.fill = CLIENT_BANNER_FILL
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20
    for c in range(1, span_cols + 1):
        ws.cell(row=row, column=c).fill = CLIENT_BANNER_FILL


CLIENT_FIELD_LABELS = [
    ("client_id", "Client ID"),
    ("legal_name", "Legal Name"),
    ("trade_name", "Trade Name"),
    ("entity_type", "Entity Type"),
    ("financial_year", "Financial Year"),
    ("accounting_framework", "Accounting Framework"),
    ("base_currency", "Base Currency"),
    ("country", "Country"),
    ("jurisdiction", "Jurisdiction"),
    ("chart_of_accounts_note", "Chart of Accounts"),
    ("tax_configuration_note", "Tax Configuration"),
    ("approval_workflow_note", "Approval Workflow"),
    ("document_storage_note", "Document Storage Area"),
    ("journal_numbering_prefix", "Journal Numbering Prefix"),
]


def build_client_info_sheet(ws, client):
    """First tab of the workbook when a client is specified -- the workbook's equivalent of
    'confirm the correct client before doing anything else'. One workbook = one client's
    workspace; this sheet is that workspace's cover page."""
    ws.title = "Client Info"
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value="Client Information").font = TITLE_FONT
    ws.cell(
        row=2, column=1,
        value="Confirm this is the correct client/company before creating, reviewing, or approving any "
              "entry in this workbook. This workbook contains records for this client ONLY.",
    ).font = SUBTITLE_FONT
    ws.merge_cells("A2:D2")
    add_client_banner(ws, 4, client, span_cols=4)

    row = 6
    for key, label in CLIENT_FIELD_LABELS:
        val = client.get(key, "")
        if not val:
            continue
        ws.cell(row=row, column=1, value=label).font = Font(name=FONT_NAME, bold=True, size=10, color="555555")
        ws.cell(row=row, column=2, value=val)
        row += 1

    reg_ids = client.get("registration_ids", {}) or {}
    for reg_key, reg_val in reg_ids.items():
        ws.cell(row=row, column=1, value=reg_key).font = Font(name=FONT_NAME, bold=True, size=10, color="555555")
        ws.cell(row=row, column=2, value=reg_val)
        row += 1

    branches = client.get("branches", []) or []
    if branches:
        ws.cell(row=row, column=1, value="Branches / Locations").font = Font(name=FONT_NAME, bold=True, size=10, color="555555")
        ws.cell(row=row, column=2, value="; ".join(branches))
        row += 1

    users = client.get("authorized_users", []) or []
    if users:
        ws.cell(row=row, column=1, value="Authorized Users").font = Font(name=FONT_NAME, bold=True, size=10, color="555555")
        ws.cell(row=row, column=2, value="; ".join(users))
        row += 1

    row += 1
    ws.cell(
        row=row, column=1,
        value="Data separation note: this workbook is scoped to this client only. Do not copy rows from another "
              "client's workbook into this one, or vice versa, outside a controlled reassignment process.",
    ).font = Font(name=FONT_NAME, italic=True, size=9, color="B00020")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

    autosize(ws, [26, 40])
    return ws


def join_unique(values):
    seen = []
    for v in values:
        if v and v not in seen:
            seen.append(v)
    return "; ".join(seen)


def account_summary(lines, side):
    """side = 'debit' or 'credit' -- join account names that have a nonzero amount on that side."""
    names = [ln.get("account_name", "") for ln in lines if (ln.get(side, 0) or 0) > 0]
    if not names:
        return ""
    if len(names) == 1:
        return names[0]
    return f"{names[0]} +{len(names) - 1} more"


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

def build_dashboard(ws, n_entries, client=None):
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False

    ws.cell(row=1, column=1, value="Journal Entry Dashboard").font = TITLE_FONT
    ws.cell(
        row=2, column=1,
        value="All figures are live formulas over the Journal Register / Journal Lines / Exceptions sheets. "
              "Use the filter arrows on those tabs to slice by date, period, status, type, vendor/customer, "
              "preparer, reviewer, department, cost centre, or project.",
    ).font = SUBTITLE_FONT
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 30
    ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    banner_row = 3
    if client:
        add_client_banner(ws, banner_row, client)

    reg = "'Journal Register'"
    lines = "'Journal Lines'"
    exc = "'Exceptions'"
    # Journal Register and Journal Lines headers are on row 3 (data starts row 4);
    # Exceptions header is on row 3 (data starts row 4). Ranges MUST start after the
    # header row or the header text gets pulled into numeric formulas -> #VALUE!.
    reg0, lines0, exc0 = 4, 4, 4

    kpis = [
        ("Total Entries", f"=COUNTA({reg}!$A${reg0}:$A${MAX_ROW})", None),
        ("Total Debit", f"=SUM({lines}!$E${lines0}:$E${MAX_ROW})", CURRENCY_FMT),
        ("Total Credit", f"=SUM({lines}!$F${lines0}:$F${MAX_ROW})", CURRENCY_FMT),
        ("Draft", f'=COUNTIF({reg}!$J${reg0}:$J${MAX_ROW},"Draft")', None),
        ("Pending Review", f'=COUNTIF({reg}!$J${reg0}:$J${MAX_ROW},"Pending Review")', None),
        ("Approved", f'=COUNTIF({reg}!$J${reg0}:$J${MAX_ROW},"Approved")', None),
        ("Posted", f'=COUNTIF({reg}!$J${reg0}:$J${MAX_ROW},"Posted")', None),
        ("Rejected", f'=COUNTIF({reg}!$J${reg0}:$J${MAX_ROW},"Rejected")', None),
        ("Reversed", f'=COUNTIF({reg}!$J${reg0}:$J${MAX_ROW},"Reversed")', None),
        ("Unbalanced Entries",
         f"=SUMPRODUCT(--(ROUND({reg}!$H${reg0}:$H${MAX_ROW}-{reg}!$I${reg0}:$I${MAX_ROW},2)<>0),"
         f"--({reg}!$A${reg0}:$A${MAX_ROW}<>\"\"))", None),
        ("Duplicate Warnings", f'=COUNTIF({exc}!$B${exc0}:$B${MAX_ROW},"*duplicate*")', None),
        ("Missing-Document Warnings",
         f'=COUNTIF({exc}!$B${exc0}:$B${MAX_ROW},"*missing*document*")',
         None),
        ("Low-Confidence AI Entries", f'=COUNTIF({reg}!$L${reg0}:$L${MAX_ROW},"Low")', None),
        ("Entries Requiring Tax Review", f'=COUNTIF({exc}!$B${exc0}:$B${MAX_ROW},"*tax*")', None),
    ]

    row, col = (5 if client else 4), 1
    cols_per_row = 4
    for i, (label, formula, fmt) in enumerate(kpis):
        r = row + (i // cols_per_row) * 3
        c = col + (i % cols_per_row) * 2
        cell_label = ws.cell(row=r, column=c, value=label)
        cell_label.font = KPI_LABEL_FONT
        cell_label.fill = KPI_FILL
        cell_value = ws.cell(row=r + 1, column=c, value=formula)
        cell_value.font = KPI_VALUE_FONT
        cell_value.fill = KPI_FILL
        if fmt:
            cell_value.number_format = fmt
        for cc in (c, c + 1):
            ws.cell(row=r, column=cc).fill = KPI_FILL
            ws.cell(row=r + 1, column=cc).fill = KPI_FILL
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)

    autosize(ws, [16] * 8)
    last_kpi_row = row + ((len(kpis) - 1) // cols_per_row) * 3 + 2
    ws.cell(
        row=last_kpi_row + 2, column=1,
        value="DRAFT workbook. No entry here has been posted by Claude -- posting only happens through your own "
              "accounting system after human review and approval.",
    ).font = Font(name=FONT_NAME, italic=True, bold=True, color="B00020")
    ws.merge_cells(start_row=last_kpi_row + 2, start_column=1, end_row=last_kpi_row + 2, end_column=8)
    return ws


# ---------------------------------------------------------------------------
# Journal Register -- one row per entry
# ---------------------------------------------------------------------------

REGISTER_HEADERS = [
    "Entry ID", "Date", "Period", "Entry Type", "Description",
    "Debit Account Summary", "Credit Account Summary", "Debit Total", "Credit Total",
    "Status", "Source", "Confidence", "Reversal Date",
    "Vendor/Customer", "Department", "Cost Centre", "Project",
    "Preparer", "Reviewer", "View Details",
    "Source Method", "Branch",
]


def build_journal_register(wb, entries, entry_detail_row, client=None):
    ws = wb.create_sheet("Journal Register")
    ws.cell(row=1, column=1, value="Journal Register").font = TITLE_FONT
    ws.merge_cells("A1:T1")
    if client:
        add_client_banner(ws, 2, client, span_cols=len(REGISTER_HEADERS))
    header_row = 3
    for c, h in enumerate(REGISTER_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(REGISTER_HEADERS))

    row = header_row + 1
    for entry in entries:
        eid = entry.get("entry_id", "")
        lines = entry.get("lines", [])
        vendor = join_unique([ln.get("vendor_customer", "") for ln in lines])
        dept = join_unique([ln.get("department", "") for ln in lines])
        cc = join_unique([ln.get("cost_centre", "") for ln in lines])
        proj = join_unique([ln.get("project", "") for ln in lines])

        ws.cell(row=row, column=1, value=eid)
        ws.cell(row=row, column=2, value=entry.get("transaction_date", ""))
        ws.cell(row=row, column=3, value=entry.get("accounting_period", ""))
        ws.cell(row=row, column=4, value=entry.get("entry_type", ""))
        ws.cell(row=row, column=5, value=entry.get("description", ""))
        ws.cell(row=row, column=6, value=account_summary(lines, "debit"))
        ws.cell(row=row, column=7, value=account_summary(lines, "credit"))

        debit_cell = ws.cell(
            row=row, column=8,
            value=f"=SUMIF('Journal Lines'!$A:$A,$A{row},'Journal Lines'!$E:$E)",
        )
        credit_cell = ws.cell(
            row=row, column=9,
            value=f"=SUMIF('Journal Lines'!$A:$A,$A{row},'Journal Lines'!$F:$F)",
        )
        debit_cell.number_format = CURRENCY_FMT
        credit_cell.number_format = CURRENCY_FMT

        status = entry.get("status", "Draft")
        status_cell = ws.cell(row=row, column=10, value=status)
        status_cell.fill = STATUS_FILL.get(status, NEUTRAL_FILL)

        ws.cell(row=row, column=11, value=entry.get("source_reference", ""))

        conf = entry.get("confidence", "")
        conf_cell = ws.cell(row=row, column=12, value=conf)
        conf_cell.fill = LEVEL_FILL.get(conf, PatternFill())

        ws.cell(row=row, column=13, value=entry.get("reversal_date", "") if entry.get("reversal_required") else "")
        ws.cell(row=row, column=14, value=vendor)
        ws.cell(row=row, column=15, value=dept)
        ws.cell(row=row, column=16, value=cc)
        ws.cell(row=row, column=17, value=proj)
        ws.cell(row=row, column=18, value=entry.get("preparer", ""))
        ws.cell(row=row, column=19, value=entry.get("reviewer", ""))

        detail_row = entry_detail_row.get(eid)
        link_cell = ws.cell(row=row, column=20, value="Open")
        if detail_row:
            link_to(link_cell, "Entry Details", f"A{detail_row}")

        branch = join_unique([ln.get("branch", "") for ln in lines])
        ws.cell(row=row, column=21, value=entry.get("source_method", ""))
        ws.cell(row=row, column=22, value=branch)

        for c in range(1, len(REGISTER_HEADERS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    n = len(entries)
    add_table(ws, "JournalRegisterTable", len(REGISTER_HEADERS), n, header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    autosize(ws, [12, 12, 10, 16, 30, 22, 22, 14, 14, 15, 16, 11, 13, 18, 14, 12, 12, 14, 14, 10, 16, 14])
    return ws


# ---------------------------------------------------------------------------
# Journal Lines -- normalized, one row per debit/credit line
# ---------------------------------------------------------------------------

LINES_HEADERS = [
    "Entry ID", "Line #", "Account Code", "Account Name", "Debit", "Credit",
    "Narration", "Vendor/Customer", "Department", "Cost Centre", "Project",
    "Tax Code", "Source Document Ref", "Branch",
]


def build_journal_lines(wb, entries, entry_detail_row, client=None):
    ws = wb.create_sheet("Journal Lines")
    ws.cell(row=1, column=1, value="Journal Lines (normalized)").font = TITLE_FONT
    if client:
        add_client_banner(ws, 2, client, span_cols=len(LINES_HEADERS))
    header_row = 3
    for c, h in enumerate(LINES_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(LINES_HEADERS))

    row = header_row + 1
    total_lines = 0
    for entry in entries:
        eid = entry.get("entry_id", "")
        detail_row = entry_detail_row.get(eid)
        for ln in entry.get("lines", []):
            id_cell = ws.cell(row=row, column=1, value=eid)
            if detail_row:
                link_to(id_cell, "Entry Details", f"A{detail_row}")
            ws.cell(row=row, column=2, value=ln.get("line_no", ""))
            ws.cell(row=row, column=3, value=ln.get("account_code", ""))
            ws.cell(row=row, column=4, value=ln.get("account_name", ""))
            debit_cell = ws.cell(row=row, column=5, value=ln.get("debit", 0) or None)
            credit_cell = ws.cell(row=row, column=6, value=ln.get("credit", 0) or None)
            debit_cell.number_format = CURRENCY_FMT
            credit_cell.number_format = CURRENCY_FMT
            ws.cell(row=row, column=7, value=ln.get("narration", ""))
            ws.cell(row=row, column=8, value=ln.get("vendor_customer", ""))
            ws.cell(row=row, column=9, value=ln.get("department", ""))
            ws.cell(row=row, column=10, value=ln.get("cost_centre", ""))
            ws.cell(row=row, column=11, value=ln.get("project", ""))
            ws.cell(row=row, column=12, value=ln.get("tax_code", ""))
            ws.cell(row=row, column=13, value=ln.get("source_doc_ref", ""))
            ws.cell(row=row, column=14, value=ln.get("branch", ""))
            for c in range(1, len(LINES_HEADERS) + 1):
                ws.cell(row=row, column=c).border = BORDER
            row += 1
            total_lines += 1

    add_table(ws, "JournalLinesTable", len(LINES_HEADERS), total_lines, header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    autosize(ws, [12, 8, 12, 26, 13, 13, 26, 18, 14, 12, 12, 10, 18, 14])
    return ws


# ---------------------------------------------------------------------------
# Entry Details -- wide table, still one row per entry, full audit trail
# ---------------------------------------------------------------------------

DETAIL_HEADERS = [
    "Entry ID", "Status", "Transaction Date", "Posting Date", "Accounting Period",
    "Debit Total", "Credit Total", "Balance Check", "Narration", "Calculation Method",
    "Assumptions", "Confidence", "Confidence Explanation", "Tax Treatment",
    "Reversal Required", "Reversal Date", "Preparer", "Reviewer", "Approver",
    "Created At", "Reviewed At", "Approved At", "Posted At", "Change History",
    "Exceptions / Warnings", "Supporting Documents", "Back to Register",
    "Source Method", "Branch", "Supporting Document Note",
]


def build_entry_details(wb, entries, exceptions_by_entry, docs_by_entry, register_row, client=None):
    ws = wb.create_sheet("Entry Details")
    ws.cell(row=1, column=1, value="Entry Details and Audit View").font = TITLE_FONT
    ws.cell(
        row=2, column=1,
        value="One row per entry -- full audit trail. Use Ctrl+F / the filter arrows to jump to an Entry ID, "
              "or click through from the Journal Register.",
    ).font = SUBTITLE_FONT
    if client:
        add_client_banner(ws, 3, client, span_cols=len(DETAIL_HEADERS))
    header_row = 4
    for c, h in enumerate(DETAIL_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(DETAIL_HEADERS))

    row = header_row + 1
    for entry in entries:
        eid = entry.get("entry_id", "")
        ws.cell(row=row, column=1, value=eid)
        status = entry.get("status", "Draft")
        status_cell = ws.cell(row=row, column=2, value=status)
        status_cell.fill = STATUS_FILL.get(status, NEUTRAL_FILL)
        ws.cell(row=row, column=3, value=entry.get("transaction_date", ""))
        ws.cell(row=row, column=4, value=entry.get("posting_date", ""))
        ws.cell(row=row, column=5, value=entry.get("accounting_period", ""))

        debit_cell = ws.cell(
            row=row, column=6,
            value=f"=SUMIF('Journal Lines'!$A:$A,$A{row},'Journal Lines'!$E:$E)",
        )
        credit_cell = ws.cell(
            row=row, column=7,
            value=f"=SUMIF('Journal Lines'!$A:$A,$A{row},'Journal Lines'!$F:$F)",
        )
        debit_cell.number_format = CURRENCY_FMT
        credit_cell.number_format = CURRENCY_FMT
        check_cell = ws.cell(
            row=row, column=8,
            value=f'=IF(ROUND(F{row}-G{row},2)=0,"Balanced","NOT BALANCED")',
        )
        check_cell.font = Font(name=FONT_NAME, bold=True)

        ws.cell(row=row, column=9, value=entry.get("narration", ""))
        ws.cell(row=row, column=10, value=entry.get("calculation_method", ""))
        ws.cell(row=row, column=11, value="; ".join(entry.get("assumptions", []) or []))

        conf = entry.get("confidence", "")
        conf_cell = ws.cell(row=row, column=12, value=conf)
        conf_cell.fill = LEVEL_FILL.get(conf, PatternFill())
        conf_details = entry.get("confidence_details", []) or []
        conf_expl = "; ".join(
            f"{d.get('field', '')} ({d.get('level', '')}): {d.get('reason', '')}".strip(": ")
            for d in conf_details if d.get("level") != "High" or d.get("reason")
        )
        ws.cell(row=row, column=13, value=conf_expl)

        ws.cell(row=row, column=14, value=entry.get("tax_treatment", ""))
        ws.cell(row=row, column=15, value="Yes" if entry.get("reversal_required") else "No")
        ws.cell(row=row, column=16, value=entry.get("reversal_date", "") if entry.get("reversal_required") else "")
        ws.cell(row=row, column=17, value=entry.get("preparer", ""))
        ws.cell(row=row, column=18, value=entry.get("reviewer", ""))
        ws.cell(row=row, column=19, value=entry.get("approver", ""))
        ws.cell(row=row, column=20, value=entry.get("created_at", ""))
        ws.cell(row=row, column=21, value=entry.get("reviewed_at", ""))
        ws.cell(row=row, column=22, value=entry.get("approved_at", ""))
        ws.cell(row=row, column=23, value=entry.get("posted_at", ""))
        ws.cell(row=row, column=24, value=" | ".join(entry.get("change_history", []) or []))

        exc_list = exceptions_by_entry.get(eid, [])
        exc_summary = "; ".join(f"{e.get('exception_type', '')} ({e.get('risk_level', '')})" for e in exc_list)
        exc_cell = ws.cell(row=row, column=25, value=exc_summary)
        if exc_list:
            exc_cell.fill = BAD_FILL if any(e.get("risk_level") == "High" for e in exc_list) else WARN_FILL

        doc_list = docs_by_entry.get(eid, [])
        doc_summary = "; ".join(
            f"{d.get('document_type', '')} {d.get('doc_number', '')} ({d.get('original_filename', '')})".strip()
            for d in doc_list
        )
        ws.cell(row=row, column=26, value=doc_summary)

        reg_row = register_row.get(eid)
        back_cell = ws.cell(row=row, column=27, value="Register")
        if reg_row:
            link_to(back_cell, "Journal Register", f"A{reg_row}")

        branch = join_unique([ln.get("branch", "") for ln in entry.get("lines", [])])
        ws.cell(row=row, column=28, value=entry.get("source_method", ""))
        ws.cell(row=row, column=29, value=branch)
        doc_missing_note = entry.get("supporting_document_missing_reason", "")
        note_cell = ws.cell(row=row, column=30, value=doc_missing_note)
        if doc_missing_note:
            note_cell.fill = WARN_FILL

        for c in range(1, len(DETAIL_HEADERS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    n = len(entries)
    add_table(ws, "EntryDetailsTable", len(DETAIL_HEADERS), n, header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    widths = [12, 15, 13, 13, 12, 13, 13, 13, 30, 16, 26, 11, 30, 20,
              10, 13, 12, 12, 12, 14, 14, 14, 14, 30, 26, 30, 11, 16, 14, 26]
    autosize(ws, widths)
    return ws


# ---------------------------------------------------------------------------
# Source Documents
# ---------------------------------------------------------------------------

DOC_HEADERS = [
    "Entry ID", "Document ID", "Document Type", "Invoice/Receipt #", "Vendor/Customer",
    "Document Date", "Document Amount", "Original Filename", "File Link",
    "OCR Confidence", "Verification Status",
]


def build_source_documents(wb, documents, entry_detail_row, client=None):
    ws = wb.create_sheet("Source Documents")
    ws.cell(row=1, column=1, value="Source Documents Index").font = TITLE_FONT
    if client:
        add_client_banner(ws, 2, client, span_cols=len(DOC_HEADERS))
    header_row = 3
    for c, h in enumerate(DOC_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(DOC_HEADERS))

    row = header_row + 1
    for doc in documents:
        eid = doc.get("entry_id", "")
        id_cell = ws.cell(row=row, column=1, value=eid)
        detail_row = entry_detail_row.get(eid)
        if detail_row:
            link_to(id_cell, "Entry Details", f"A{detail_row}")
        ws.cell(row=row, column=2, value=doc.get("document_id", ""))
        ws.cell(row=row, column=3, value=doc.get("document_type", ""))
        ws.cell(row=row, column=4, value=doc.get("doc_number", ""))
        ws.cell(row=row, column=5, value=doc.get("vendor_customer", ""))
        ws.cell(row=row, column=6, value=doc.get("document_date", ""))
        amt_cell = ws.cell(row=row, column=7, value=doc.get("document_amount", None))
        amt_cell.number_format = CURRENCY_FMT
        ws.cell(row=row, column=8, value=doc.get("original_filename", ""))
        link = doc.get("file_link", "")
        link_cell = ws.cell(row=row, column=9, value=link or "")
        if link:
            link_cell.hyperlink = link
            link_cell.font = HYPERLINK_FONT
        conf = doc.get("ocr_confidence", "")
        conf_cell = ws.cell(row=row, column=10, value=conf)
        conf_cell.fill = LEVEL_FILL.get(conf, PatternFill())
        ws.cell(row=row, column=11, value=doc.get("verification_status", ""))
        for c in range(1, len(DOC_HEADERS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    add_table(ws, "SourceDocumentsTable", len(DOC_HEADERS), len(documents), header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    autosize(ws, [12, 14, 14, 16, 18, 13, 14, 24, 20, 13, 16])
    return ws


# ---------------------------------------------------------------------------
# Exceptions
# ---------------------------------------------------------------------------

EXC_HEADERS = [
    "Entry ID", "Exception Type", "Risk Level", "Description", "Required Action",
    "Assigned Reviewer", "Resolution Status", "Resolution Notes",
]


def build_exceptions(wb, exceptions, entry_detail_row, client=None):
    ws = wb.create_sheet("Exceptions")
    ws.cell(row=1, column=1, value="Exceptions").font = TITLE_FONT
    if client:
        add_client_banner(ws, 2, client, span_cols=len(EXC_HEADERS))
    header_row = 3
    for c, h in enumerate(EXC_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(EXC_HEADERS))

    row = header_row + 1
    for exc in exceptions:
        eid = exc.get("entry_id", "")
        id_cell = ws.cell(row=row, column=1, value=eid)
        detail_row = entry_detail_row.get(eid)
        if detail_row:
            link_to(id_cell, "Entry Details", f"A{detail_row}")
        ws.cell(row=row, column=2, value=exc.get("exception_type", ""))
        risk = exc.get("risk_level", "")
        risk_cell = ws.cell(row=row, column=3, value=risk)
        risk_cell.fill = LEVEL_FILL.get(risk, PatternFill())
        ws.cell(row=row, column=4, value=exc.get("description", ""))
        ws.cell(row=row, column=5, value=exc.get("required_action", ""))
        ws.cell(row=row, column=6, value=exc.get("assigned_reviewer", ""))
        ws.cell(row=row, column=7, value=exc.get("resolution_status", "Open"))
        ws.cell(row=row, column=8, value=exc.get("resolution_notes", ""))
        for c in range(1, len(EXC_HEADERS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    add_table(ws, "ExceptionsTable", len(EXC_HEADERS), len(exceptions), header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    autosize(ws, [12, 20, 11, 40, 30, 16, 16, 30])
    return ws


# ---------------------------------------------------------------------------
# Review and Approval Status
# ---------------------------------------------------------------------------

APPROVAL_HEADERS = [
    "Entry ID", "Current Status", "Preparer", "Reviewer", "Approver",
    "Prepared Date", "Reviewed Date", "Approved Date", "Posting Date",
    "Rejection Reason", "Correction/Reversal Reference",
]


def build_approval_status(wb, entries, entry_detail_row, client=None):
    ws = wb.create_sheet("Review and Approval Status")
    ws.cell(row=1, column=1, value="Review and Approval Tracker").font = TITLE_FONT
    if client:
        add_client_banner(ws, 2, client, span_cols=len(APPROVAL_HEADERS))
    header_row = 3
    for c, h in enumerate(APPROVAL_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(APPROVAL_HEADERS))

    row = header_row + 1
    for entry in entries:
        eid = entry.get("entry_id", "")
        id_cell = ws.cell(row=row, column=1, value=eid)
        detail_row = entry_detail_row.get(eid)
        if detail_row:
            link_to(id_cell, "Entry Details", f"A{detail_row}")
        status = entry.get("status", "Draft")
        status_cell = ws.cell(row=row, column=2, value=status)
        status_cell.fill = STATUS_FILL.get(status, NEUTRAL_FILL)
        ws.cell(row=row, column=3, value=entry.get("preparer", ""))
        ws.cell(row=row, column=4, value=entry.get("reviewer", ""))
        ws.cell(row=row, column=5, value=entry.get("approver", ""))
        ws.cell(row=row, column=6, value=entry.get("created_at", ""))
        ws.cell(row=row, column=7, value=entry.get("reviewed_at", ""))
        ws.cell(row=row, column=8, value=entry.get("approved_at", ""))
        ws.cell(row=row, column=9, value=entry.get("posted_at", ""))
        ws.cell(row=row, column=10, value=entry.get("rejection_reason", ""))
        ws.cell(row=row, column=11, value=entry.get("correction_reference", ""))
        for c in range(1, len(APPROVAL_HEADERS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1

    add_table(ws, "ApprovalStatusTable", len(APPROVAL_HEADERS), len(entries), header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    autosize(ws, [12, 18, 14, 14, 14, 13, 13, 13, 13, 26, 26])
    return ws


# ---------------------------------------------------------------------------
# Whole-workbook assembly (used for both single-client and multi-client runs)
# ---------------------------------------------------------------------------

def build_workbook(entries, exceptions, documents, client=None):
    if not entries:
        raise SystemExit("No entries found for this client/workbook.")

    # Row maps computed up front (independent of sheet build order).
    entry_detail_row = {e.get("entry_id", ""): 5 + i for i, e in enumerate(entries)}
    register_row = {e.get("entry_id", ""): 4 + i for i, e in enumerate(entries)}

    exceptions_by_entry = {}
    for exc in exceptions:
        exceptions_by_entry.setdefault(exc.get("entry_id", ""), []).append(exc)
    docs_by_entry = {}
    for doc in documents:
        docs_by_entry.setdefault(doc.get("entry_id", ""), []).append(doc)

    wb = Workbook()
    if client:
        info_ws = wb.active
        build_client_info_sheet(info_ws, client)
        dash_ws = wb.create_sheet("Dashboard")
    else:
        dash_ws = wb.active
    build_dashboard(dash_ws, len(entries), client=client)
    build_journal_register(wb, entries, entry_detail_row, client=client)
    build_journal_lines(wb, entries, entry_detail_row, client=client)
    build_entry_details(wb, entries, exceptions_by_entry, docs_by_entry, register_row, client=client)
    build_source_documents(wb, documents, entry_detail_row, client=client)
    build_exceptions(wb, exceptions, entry_detail_row, client=client)
    build_approval_status(wb, entries, entry_detail_row, client=client)

    wb.active = 0  # open on Client Info (if present) or Dashboard
    return wb


def client_kpis(entries, exceptions):
    """Plain-Python KPI counts for one client -- used by the portfolio dashboard, which
    summarizes across separate closed workbooks rather than relying on fragile cross-file
    Excel formulas."""
    status_counts = {}
    for e in entries:
        s = e.get("status", "Draft")
        status_counts[s] = status_counts.get(s, 0) + 1
    total_debit = sum(ln.get("debit", 0) or 0 for e in entries for ln in e.get("lines", []))
    total_credit = sum(ln.get("credit", 0) or 0 for e in entries for ln in e.get("lines", []))
    unbalanced = 0
    for e in entries:
        d = sum(ln.get("debit", 0) or 0 for ln in e.get("lines", []))
        c = sum(ln.get("credit", 0) or 0 for ln in e.get("lines", []))
        if round(d - c, 2) != 0:
            unbalanced += 1
    dup = sum(1 for x in exceptions if "duplicate" in x.get("exception_type", "").lower())
    missing_doc = sum(1 for x in exceptions if "missing" in x.get("exception_type", "").lower()
                       and "document" in x.get("exception_type", "").lower())
    tax = sum(1 for x in exceptions if "tax" in x.get("exception_type", "").lower())
    low_conf = sum(1 for e in entries if e.get("confidence") == "Low")
    return {
        "total_entries": len(entries),
        "total_debit": total_debit,
        "total_credit": total_credit,
        "draft": status_counts.get("Draft", 0),
        "pending_review": status_counts.get("Pending Review", 0),
        "approved": status_counts.get("Approved", 0),
        "posted": status_counts.get("Posted", 0),
        "rejected": status_counts.get("Rejected", 0),
        "unbalanced": unbalanced,
        "duplicate_warnings": dup,
        "missing_doc_warnings": missing_doc,
        "tax_review": tax,
        "low_confidence": low_conf,
    }


PORTFOLIO_HEADERS = [
    "Client ID", "Client Name", "Total Entries", "Total Debit", "Total Credit",
    "Draft", "Pending Review", "Approved", "Posted", "Rejected", "Unbalanced",
    "Duplicate Warnings", "Missing-Doc Warnings", "Tax Review Needed", "Low-Confidence",
]


def build_portfolio_dashboard(client_summaries):
    """A CA-firm / portfolio-level view across clients. Summarized counts only, computed at
    build time in Python -- each client's own workbook remains the source of full detail.
    Respect the user's own access permissions when deciding who this file is shared with;
    Claude has no way to enforce that itself."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio Dashboard"
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value="Portfolio Dashboard -- All Clients").font = TITLE_FONT
    ws.cell(
        row=2, column=1,
        value="Summarized information only. Open each client's own workbook for full entry-level detail "
              "and audit trail. Share this file only with users authorized to see cross-client summaries.",
    ).font = SUBTITLE_FONT
    ws.merge_cells("A2:F2")

    header_row = 4
    for c, h in enumerate(PORTFOLIO_HEADERS, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(PORTFOLIO_HEADERS))

    row = header_row + 1
    first_data_row = row
    for cid, cname, kpi in client_summaries:
        ws.cell(row=row, column=1, value=cid)
        ws.cell(row=row, column=2, value=cname)
        ws.cell(row=row, column=3, value=kpi["total_entries"])
        ws.cell(row=row, column=4, value=kpi["total_debit"]).number_format = CURRENCY_FMT
        ws.cell(row=row, column=5, value=kpi["total_credit"]).number_format = CURRENCY_FMT
        ws.cell(row=row, column=6, value=kpi["draft"])
        ws.cell(row=row, column=7, value=kpi["pending_review"])
        ws.cell(row=row, column=8, value=kpi["approved"])
        ws.cell(row=row, column=9, value=kpi["posted"])
        ws.cell(row=row, column=10, value=kpi["rejected"])
        unb_cell = ws.cell(row=row, column=11, value=kpi["unbalanced"])
        if kpi["unbalanced"]:
            unb_cell.fill = BAD_FILL
        ws.cell(row=row, column=12, value=kpi["duplicate_warnings"])
        ws.cell(row=row, column=13, value=kpi["missing_doc_warnings"])
        ws.cell(row=row, column=14, value=kpi["tax_review"])
        ws.cell(row=row, column=15, value=kpi["low_confidence"])
        for c in range(1, len(PORTFOLIO_HEADERS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        row += 1
    last_data_row = row - 1

    total_row = row + 1
    ws.cell(row=total_row, column=2, value="Total").font = BOLD_FONT
    for col in (3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
        cell.font = BOLD_FONT
        if col in (4, 5):
            cell.number_format = CURRENCY_FMT

    add_table(ws, "PortfolioTable", len(PORTFOLIO_HEADERS), len(client_summaries), header_row=header_row)
    ws.freeze_panes = f"B{header_row + 1}"
    autosize(ws, [12, 26, 12, 14, 14, 9, 13, 10, 9, 10, 11, 12, 14, 14, 13])
    return wb


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", required=True, help="Path to input JSON file (see module docstring for shape)")
    parser.add_argument("--output", required=True,
                         help="Single-client mode: path to write the .xlsx. "
                              "Multi-client mode (top-level 'clients' key in input): a directory to write "
                              "one workbook per client plus portfolio_dashboard.xlsx into.")
    args = parser.parse_args()

    with open(args.input, "r") as f:
        data = json.load(f)

    if "clients" in data:
        # ---- Multi-client mode ----
        out_dir = args.output
        os.makedirs(out_dir, exist_ok=True)
        client_summaries = []
        written = []
        for block in data["clients"]:
            client = block.get("client", {})
            cid = client.get("client_id", "UNKNOWN")
            entries = block.get("entries", [])
            exceptions = block.get("exceptions", [])
            documents = block.get("source_documents", [])
            if not entries:
                print(f"Skipping client {cid}: no entries.")
                continue
            wb = build_workbook(entries, exceptions, documents, client=client)
            safe_cid = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in cid)
            out_path = os.path.join(out_dir, f"{safe_cid}_journal_entries.xlsx")
            wb.save(out_path)
            written.append(out_path)
            client_summaries.append((cid, client.get("legal_name", ""), client_kpis(entries, exceptions)))
            print(f"Wrote {out_path} ({len(entries)} entries) for {cid} - {client.get('legal_name','')}")

        portfolio_wb = build_portfolio_dashboard(client_summaries)
        portfolio_path = os.path.join(out_dir, "portfolio_dashboard.xlsx")
        portfolio_wb.save(portfolio_path)
        written.append(portfolio_path)
        print(f"Wrote {portfolio_path} summarizing {len(client_summaries)} client(s).")
        print("Remember to run recalc.py on EACH file before presenting it:")
        for p in written:
            print("  python /mnt/skills/public/xlsx/scripts/recalc.py " + p)
        return

    # ---- Single-client / single-company mode ----
    entries = data.get("entries", [])
    exceptions = data.get("exceptions", [])
    documents = data.get("source_documents", [])
    client = data.get("client")  # optional; omit entirely for plain single-company use

    wb = build_workbook(entries, exceptions, documents, client=client)
    wb.save(args.output)
    print(f"Wrote {args.output} with {len(entries)} entr{'y' if len(entries) == 1 else 'ies'}, "
          f"{len(exceptions)} exception(s), {len(documents)} source document(s)."
          + (f" Client: {client.get('client_id','')} - {client.get('legal_name','')}" if client else ""))
    print("Remember to run recalc.py on this file before presenting it, so formulas have cached values:")
    print("  python /mnt/skills/public/xlsx/scripts/recalc.py " + args.output)


EXAMPLE_JSON = {
    "client": {
        "client_id": "CLI-0001",
        "legal_name": "ABC Private Limited",
        "trade_name": "",
        "entity_type": "Private Limited Company",
        "financial_year": "2026-27",
        "accounting_framework": "Ind AS",
        "base_currency": "INR",
        "country": "India",
        "jurisdiction": "India",
        "registration_ids": {"PAN": "ABCDE1234F", "GSTIN": "27ABCDE1234F1Z5"},
        "branches": ["HQ - Mumbai", "Branch - Pune"],
        "chart_of_accounts_note": "Firm-standard COA v3",
        "tax_configuration_note": "GST monthly filer",
        "approval_workflow_note": "Two-level: Manager then Partner",
        "authorized_users": ["J. Rao (Preparer)", "P. Shah (Approver)"],
        "document_storage_note": "Firm DMS, /Clients/CLI-0001/",
        "journal_numbering_prefix": "CLI-0001-JE-2026-",
    },
    "entries": [
        {
            "entry_id": "CLI-0001-JE-2026-00001",
            "status": "Pending Review",
            "entry_type": "Purchase",
            "source_method": "Invoice Upload",
            "transaction_date": "2026-08-31",
            "posting_date": "2026-08-31",
            "accounting_period": "Aug-2026",
            "description": "Office supplies - Vendor X",
            "source_reference": "Invoice 145",
            "currency": "INR",
            "narration": "Being office supplies purchased from Vendor X under Invoice No. 145, dated 31 Aug 2026.",
            "calculation_method": "Direct from invoice",
            "assumptions": ["Assumed accrual basis per prior entries."],
            "confidence": "Medium",
            "confidence_details": [
                {"field": "Account classification", "level": "Medium", "reason": "Vendor not seen before."}
            ],
            "tax_treatment": "18% GST (9% CGST + 9% SGST), intra-state",
            "reversal_required": False,
            "reversal_date": "",
            "preparer": "AI Draft",
            "reviewer": "",
            "approver": "",
            "created_at": "2026-09-01 10:00",
            "reviewed_at": "",
            "approved_at": "",
            "posted_at": "",
            "change_history": [],
            "rejection_reason": "",
            "correction_reference": "",
            "supporting_document_missing_reason": "",
            "lines": [
                {"line_no": 1, "account_code": "5010", "account_name": "Office Supplies Expense",
                 "debit": 10000, "credit": 0, "narration": "Supplies", "vendor_customer": "Vendor X",
                 "department": "Admin", "cost_centre": "HQ", "project": "", "tax_code": "GST18",
                 "source_doc_ref": "Invoice 145", "branch": "HQ - Mumbai"},
                {"line_no": 2, "account_code": "1360", "account_name": "Input GST",
                 "debit": 1800, "credit": 0, "narration": "Input GST", "vendor_customer": "Vendor X",
                 "department": "Admin", "cost_centre": "HQ", "project": "", "tax_code": "GST18",
                 "source_doc_ref": "Invoice 145", "branch": "HQ - Mumbai"},
                {"line_no": 3, "account_code": "2010", "account_name": "Accounts Payable - Vendor X",
                 "debit": 0, "credit": 11800, "narration": "Vendor payable", "vendor_customer": "Vendor X",
                 "department": "Admin", "cost_centre": "HQ", "project": "", "tax_code": "",
                 "source_doc_ref": "Invoice 145", "branch": "HQ - Mumbai"},
            ],
        }
    ],
    "exceptions": [
        {"entry_id": "CLI-0001-JE-2026-00001", "exception_type": "Low-confidence classification", "risk_level": "Medium",
         "description": "Vendor X not previously seen; expense category inferred from item description.",
         "required_action": "Confirm expense account with preparer.", "assigned_reviewer": "",
         "resolution_status": "Open", "resolution_notes": ""}
    ],
    "source_documents": [
        {"entry_id": "CLI-0001-JE-2026-00001", "document_id": "DOC-0001", "document_type": "Invoice", "doc_number": "145",
         "vendor_customer": "Vendor X", "document_date": "2026-08-31", "document_amount": 11800,
         "original_filename": "vendor_x_invoice_145.pdf", "file_link": "", "ocr_confidence": "High",
         "verification_status": "Not yet verified"}
    ],
}

# Multi-client shape: top-level {"clients": [{"client": {...}, "entries": [...],
# "exceptions": [...], "source_documents": [...]}, ...]}. Each block is exactly the shape
# of EXAMPLE_JSON above (minus needing to repeat "clients"). Run with --output pointing at a
# DIRECTORY in this mode -- the script writes one workbook per client plus
# portfolio_dashboard.xlsx.

if __name__ == "__main__":
    main()
